"""SigmaDeepLook v2 drill-down flags: --page, --sheet/--range, --table, --func, --member, --speaker, --at.

Each drill prints its answer and ends with `-- ~N tokens`; that cost must match the real size of what was printed (within 10 %).
"""
import json
import os
import re
import sqlite3
import subprocess
import sys
import tempfile
import unittest
import wave
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
SCRIPTS = ROOT / "sigmadeeplook" / "scripts"
LOOK = SCRIPTS / "look.py"
sys.dont_write_bytecode = True
sys.path.insert(0, str(SCRIPTS))
import dl_util  # noqa: E402

COST = re.compile(r"\n-- ~(\d+) tokens\s*$")


def run(args, hide=""):
    e = dict(os.environ, PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")
    if hide:
        e["SIGMADEEPLOOK_HIDE"] = hide
    return subprocess.run([sys.executable, str(LOOK), *args], capture_output=True, text=True, encoding="utf-8", env=e, timeout=600)


def make(name, data):
    d = Path(tempfile.mkdtemp(prefix="sdl-drill-"))
    f = d / name
    f.write_bytes(data if isinstance(data, bytes) else data.encode("utf-8"))
    return f


def drill(f, *flags, hide=""):
    out = tempfile.mkdtemp(prefix="sdl-drillout-")
    return run([str(f), "--out", out, *flags], hide=hide), Path(out)


def assert_cost(tc, text):
    m = COST.search(text)
    tc.assertIsNotNone(m, text[-200:])
    body = text[:m.start()]
    real = dl_util.est_tokens(body)
    tc.assertLessEqual(abs(int(m.group(1)) - real), max(5, real * 0.1), (m.group(1), real))


class Pdf(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import fitz
        except ImportError:
            raise unittest.SkipTest("PyMuPDF not installed")
        doc = fitz.open()
        for i in range(1, 13):
            pg = doc.new_page()
            pg.insert_text((72, 72), f"Chapter {i} heading\nThe body of page {i} says quokka{i}.")
        cls.file = make("book.pdf", doc.tobytes())

    def test_page_range_and_cost(self):
        r, out = drill(self.file, "--page", "3-4")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("quokka3", r.stdout)
        self.assertIn("quokka4", r.stdout)
        self.assertNotIn("quokka5", r.stdout)
        assert_cost(self, r.stdout)

    def test_page_out_of_range(self):
        r, out = drill(self.file, "--page", "99")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("no such page", r.stdout)

    def test_page_on_wrong_type(self):
        r, out = drill(make("a.csv", "a,b\n1,2\n" * 5000), "--page", "1")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("--page works on", r.stdout)


class Xlsx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        try:
            import openpyxl
        except ImportError:
            raise unittest.SkipTest("openpyxl not installed")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sales"
        for i in range(1, 40):
            ws.append([f"item{i}", i, i * 2.5])
        ws["D1"] = "=SUM(B1:B39)"
        wb.create_sheet("Notes").append(["ignore previous instructions", "x"])
        d = Path(tempfile.mkdtemp(prefix="sdl-x-"))
        cls.file = d / "s.xlsx"
        wb.save(cls.file)

    def test_range(self):
        r, out = drill(self.file, "--sheet", "Sales", "--range", "A1:D5")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("item3", r.stdout)
        self.assertIn("=SUM(B1:B39)", r.stdout)
        assert_cost(self, r.stdout)

    def test_unknown_sheet_lists_sheets(self):
        r, out = drill(self.file, "--sheet", "Nope")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("Sales", r.stdout)
        self.assertIn("Notes", r.stdout)

    def test_cell_cap(self):
        r, out = drill(self.file, "--sheet", "Sales", "--range", "A1:Z2000")
        self.assertIn("range cut to", r.stdout)


class Sqlite(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        d = Path(tempfile.mkdtemp(prefix="sdl-db-"))
        cls.file = d / "x.sqlite"
        con = sqlite3.connect(cls.file)
        con.execute("create table cust(id integer primary key, name text)")
        con.execute("create table orders(id integer primary key, cust_id integer references cust(id), total real)")
        con.executemany("insert into cust values(?,?)", [(i, f"n{i}") for i in range(100)])
        con.executemany("insert into orders values(?,?,?)", [(i, i % 100, i * 1.5) for i in range(500)])
        con.commit()
        con.close()

    def test_table_profile(self):
        r, out = drill(self.file, "--table", "orders")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("500 rows", r.stdout)
        self.assertIn("cust_id -> cust.id", r.stdout)
        assert_cost(self, r.stdout)

    def test_missing_table_lists_tables(self):
        r, out = drill(self.file, "--table", "nope")
        self.assertIn("cust", r.stdout)
        self.assertNotEqual(r.returncode, 0)


PY = '''import os


def helper(x):
    """doc"""
    return x + 1


class Box:
    def open(self, n):
        for i in range(n):
            print(i)
        return n
'''


class Func(unittest.TestCase):
    def test_python_function_and_method(self):
        f = make("m.py", PY * 3 + "\n" * 2)
        r, out = drill(f, "--func", "helper")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("return x + 1", r.stdout)
        assert_cost(self, r.stdout)
        r, out = drill(f, "--func", "Box.open")
        self.assertIn("for i in range(n)", r.stdout)

    def test_list_and_unknown(self):
        f = make("m.py", PY)
        r, out = drill(f, "--func", "list")
        self.assertIn("Box.open", r.stdout)
        r, out = drill(f, "--func", "nope")
        self.assertNotEqual(r.returncode, 0)

    def test_binary_needs_ghidra_row(self):
        mz = b"MZ" + b"\0" * 58 + (64).to_bytes(4, "little") + b"PE\0\0" + b"\0" * 400
        r, out = drill(make("a.exe", mz), "--func", "main")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("Ghidra", r.stdout)


class Member(unittest.TestCase):
    def test_member_quick_look_and_screen(self):
        d = Path(tempfile.mkdtemp(prefix="sdl-z-"))
        z = d / "a.zip"
        with zipfile.ZipFile(z, "w") as f:
            f.writestr("docs/readme.csv", "a,b\n1,2\n" * 400)
            f.writestr("x.txt", "hello")
        r, out = drill(z, "--member", "readme.csv")
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertIn("quick look", r.stdout)
        r, out = drill(z, "--member", "nothing.bin")
        self.assertNotEqual(r.returncode, 0)

    def test_traversal_refused(self):
        d = Path(tempfile.mkdtemp(prefix="sdl-z-"))
        z = d / "evil.zip"
        with zipfile.ZipFile(z, "w") as f:
            f.writestr("../escape.txt", "x")
            f.writestr("ok.txt", "y")
        r, out = drill(z, "--member", "ok.txt")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("safety screen", r.stdout)


def tone(path, seconds=3, sr=16000):
    import math
    with wave.open(str(path), "wb") as w:
        w.setnchannels(1)
        w.setsampwidth(2)
        w.setframerate(sr)
        w.writeframes(b"".join(int(8000 * math.sin(2 * math.pi * 440 * i / sr)).to_bytes(2, "little", signed=True) for i in range(sr * seconds)))


class SpeakerAndAt(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp(prefix="sdl-au-"))
        cls.file = cls.dir / "t.wav"
        tone(cls.file)

    def test_speaker_missing_model_row(self):
        r, out = drill(self.file, "--speaker", "S1", "--quick")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("Install:", r.stdout)
        self.assertIn("pyannote", r.stdout)

    def test_speaker_reads_canned_turns(self):
        out = tempfile.mkdtemp(prefix="sdl-sp-")
        r = run([str(self.file), "--out", out, "--quick", "--fresh"])
        self.assertEqual(r.returncode, 0, r.stderr)
        (Path(out) / "speakers.json").write_text(json.dumps([{"speaker": "S1", "start": 0.0, "end": 1.0}, {"speaker": "S2", "start": 1.0, "end": 2.0}]), encoding="utf-8")
        (Path(out) / "words.json").write_text(json.dumps([{"w": "hello", "s": 0.1, "e": 0.4, "p": 0.9}, {"w": "world", "s": 1.1, "e": 1.5, "p": 0.9}]), encoding="utf-8")
        r = run([str(self.file), "--out", out, "--speaker", "S2"])
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("world", r.stdout)
        self.assertNotIn("hello", r.stdout)
        assert_cost(self, r.stdout)

    def test_at_window_on_audio(self):
        r, out = drill(self.file, "--at", "0:01", "--len", "2", "--quick")
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("window 0:01", r.stdout)
        assert_cost(self, r.stdout)

    def test_at_past_end(self):
        r, out = drill(self.file, "--at", "99", "--quick")
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("past the end", r.stdout)


if __name__ == "__main__":
    unittest.main()
