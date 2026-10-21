"""The report as a 16:9 deck (python-pptx), Ledger look: Calibri, navy rules, the answer on the first slide.

Charts are native PowerPoint charts, so the numbers stay editable: the waterfall is a stacked column with an invisible base
series, the football field a stacked bar with a transparent first series. Body text 18 pt, chart labels 12 pt or more.
"""
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

import fin_svg as S

RGB = {k: RGBColor.from_string(v[1:]) for k, v in S.TOK.items()}
FONT = "Calibri"
NUMFMT = {"money": '#,##0_);(#,##0)', "pct": "0.0%", "x": '0.0"x"', "ps": '$#,##0.00_);($#,##0.00)', "int": "0", "num1": "0.0"}


def _text(slide, x, y, w, h, text, size=18, bold=False, color="ink", align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    if align:
        p.alignment = align
    for r in p.runs:
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(size), bold, RGB[color]
    return box


def _rule(slide, y):
    ln = slide.shapes.add_shape(1, Inches(0.6), Inches(y), Inches(12.13), Inches(0.04))
    ln.fill.solid()
    ln.fill.fore_color.rgb = RGB["navy"]
    ln.line.fill.background()


def _title(slide, text, sub=None):
    _text(slide, 0.6, 0.35, 12.1, 0.7, text, 28, True, "navy")
    _rule(slide, 1.05)
    if sub:
        _text(slide, 0.6, 1.12, 12.1, 0.5, sub, 14, color="muted")


def _style_chart(chart, legend=True, fmt="money"):
    chart.font.name, chart.font.size = FONT, Pt(12)
    chart.has_title = False   # the slide title names the chart
    try:
        chart.value_axis.tick_labels.number_format = NUMFMT.get(fmt, NUMFMT["money"])
        chart.value_axis.tick_labels.number_format_is_linked = False
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = RGB["rule"]
    except (AttributeError, ValueError):
        pass
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)


def _labels(series, fmt):
    series.data_labels.show_value = True
    series.data_labels.number_format = NUMFMT.get(fmt, NUMFMT["money"])
    series.data_labels.number_format_is_linked = False
    series.data_labels.font.size = Pt(12)
    series.data_labels.font.name = FONT


def chart_slide(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, c["title"])
    x, y, w, h = Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.6)
    fmt = c.get("fmt", "money")
    if c["kind"] == "waterfall":
        run, base, vis, colors = 0.0, [], [], []
        for label, v, kind in c["steps"]:
            if kind == "total":
                base.append(0.0), vis.append(v)
                run = v
            else:
                lo = min(run, run + v)
                base.append(lo), vis.append(abs(v))
                run += v
            colors.append("navy" if kind == "total" else "steel" if kind == "part" else ("gain" if v >= 0 else "loss"))
        cd = CategoryChartData()
        cd.categories = [st[0] + chr(10) + S.fmt(st[1], fmt) for st in c["steps"]]   # the value under each bar: small bars stay readable
        cd.add_series("base", base)
        cd.add_series("value", vis)
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, x, y, w, h, cd).chart
        _style_chart(ch, False, fmt)
        ch.plots[0].gap_width = 60
        ch.series[0].format.fill.background()
        ch.series[0].format.line.fill.background()
        for i, col in enumerate(colors):
            pt = ch.series[1].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGB[col]
    elif c["kind"] == "bars":
        cd = CategoryChartData()
        cd.categories = c["labels"]
        for name, vals, _ in c["series"]:
            cd.add_series(name, [v if v is not None else 0 for v in vals])
        ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, cd).chart
        _style_chart(ch, len(c["series"]) > 1, fmt)
        for ser, (_, _, role) in zip(ch.series, c["series"]):
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = RGB["navy" if role == "total" else "steel" if role == "part" else "muted"]
            if len(c["labels"]) * len(c["series"]) <= (13 if len(c["series"]) == 1 else 12):
                _labels(ser, fmt)
    elif c["kind"] in ("hbars", "football"):
        if c["kind"] == "football":
            rows = list(c["rows"]) + [(f"{c['marker'][0]} {S.F.per_share(c['marker'][1])}", c["marker"][1], c["marker"][1] * 1.004)]
            rows = [r for r in rows if r[1] is not None and r[2] is not None]
            base, vis, cats = [r[1] for r in rows], [r[2] - r[1] for r in rows], [r[0] for r in rows]
            fmt = "ps"
        else:
            cats, vis = [r[0] for r in c["rows"]], [r[1] or 0 for r in c["rows"]]
            base = None
        cd = CategoryChartData()
        cd.categories = list(reversed(cats))
        if base is not None:
            cd.add_series("low", list(reversed(base)))
        cd.add_series("value", list(reversed(vis)))
        kind = XL_CHART_TYPE.BAR_STACKED if base is not None else XL_CHART_TYPE.BAR_CLUSTERED
        ch = s.shapes.add_chart(kind, x, y, w, h, cd).chart
        _style_chart(ch, False, fmt)
        vs = ch.series[-1]
        if base is not None:
            ch.series[0].format.fill.background()
            ch.series[0].format.line.fill.background()
            ch.value_axis.minimum_scale = round(min(base) * 0.9, 0)
        n = len(cats)
        for i in range(n):
            src = n - 1 - i
            pt = vs.points[i]
            pt.format.fill.solid()
            if c["kind"] == "football":
                col = "marker" if src == n - 1 else "steel"
            elif c.get("status"):
                col = "gain" if c["rows"][src][1] else "loss"
            else:
                col = "navy" if c["rows"][src][0] == c.get("highlight") else "steel"
            pt.format.fill.fore_color.rgb = RGB[col]
        if c["kind"] == "hbars" and not c.get("status"):
            _labels(vs, fmt)
        ch.plots[0].gap_width = 50
    return s


def table_slide(prs, t, max_rows=14):
    rows = t["rows"]
    parts = [rows[i:i + max_rows] for i in range(0, len(rows), max_rows)] or [[]]
    for pi, part in enumerate(parts):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        _title(s, t["title"] + (" (continued)" if pi else ""))
        ncol = len(t["columns"])
        first = 3.4 if ncol <= 7 else 1.6
        widths = [first] + [(12.1 - first) / max(1, ncol - 1)] * (ncol - 1) if ncol > 1 else [12.1]
        if t["formats"][-1] == "text" and ncol == 3:   # Input | Value | Source
            widths = [2.6, 1.8, 7.7]
        shape = s.shapes.add_table(len(part) + 1, ncol, Inches(0.6), Inches(1.45), Inches(12.1), Inches(0.32 * (len(part) + 1)))
        tbl = shape.table
        for ci, wdt in enumerate(widths):
            tbl.columns[ci].width = Inches(wdt)
        for ci, name in enumerate(t["columns"]):
            _cell(tbl.cell(0, ci), name, True, "bg", "navy", PP_ALIGN.LEFT if (ci == 0 or t["formats"][ci] == "text") else PP_ALIGN.RIGHT)
        for ri, row in enumerate(part):
            gi = pi * max_rows + ri
            for ci, v in enumerate(row):
                kind = t["formats"][ci]
                if ci > 0 and t.get("row_formats", {}).get(gi) and kind in ("money", "auto"):
                    kind = t["row_formats"][gi]
                txt = S.cell(v, kind)[0].replace('<span class="pad">)</span>', "").replace("&nbsp;", "")
                txt = S.html.unescape(txt)
                bold = gi in t["totals"] or gi in t["double"] or (t.get("center") == (gi, ci))
                color = "ink"
                if t.get("status_col") == ci:
                    color = "gain" if str(v).lower() in ("pass", "clear") else "loss"
                fill = "center" if t.get("center") == (gi, ci) else ("panel" if gi in t["totals"] or gi in t["double"] else "bg")
                _cell(tbl.cell(ri + 1, ci), txt, bold, color, fill, PP_ALIGN.LEFT if (ci == 0 or kind == "text") else PP_ALIGN.RIGHT)
    return s


def _cell(c, text, bold, color, fill, align):
    c.fill.solid()
    c.fill.fore_color.rgb = RGB[fill]
    c.margin_left = c.margin_right = Inches(0.06)
    c.margin_top = c.margin_bottom = Inches(0.02)
    p = c.text_frame.paragraphs[0]
    p.text = str(text)
    p.alignment = align
    for r in p.runs:
        r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(12), bold, RGB[color]


def deck(rep, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, rep["title"], rep["subtitle"])
    _text(s, 0.6, 1.9, 12, 1.2, rep["hero"]["value"], 60, True, "navy")
    _text(s, 0.6, 3.1, 12, 0.5, rep["hero"]["label"], 18, color="muted")
    _text(s, 0.6, 3.7, 12, 1.0, rep["hero"]["line"], 20)
    kp = rep.get("kpis", [])[:5]
    for i, (k, v) in enumerate(kp):
        x = 0.6 + i * (12.1 / max(1, len(kp)))
        _text(s, x, 5.3, 2.4, 0.4, k, 14, color="muted")
        _text(s, x, 5.7, 2.4, 0.6, v, 22, True)
    for c in rep.get("charts", []):
        chart_slide(prs, c)
    for t in rep.get("tables", []):
        table_slide(prs, t)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _title(s, "Notes and sources")
    lines = [f"- {n}" for n in rep.get("notes", [])] + [""] + [f"Source: {x}" for x in rep.get("sources", []) if x]
    box = _text(s, 0.6, 1.5, 12.1, 5.5, lines[0] if lines else "", 18)
    for ln in lines[1:]:
        p = box.text_frame.add_paragraph()
        p.text = ln
        for r in p.runs:
            r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(16 if ln.startswith("Source") else 18), RGB["ink" if not ln.startswith("Source") else "muted"]
    prs.save(path)
    return path
