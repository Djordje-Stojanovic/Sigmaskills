#!/usr/bin/env python3
"""sigmafinanceman: finance models from public filings, as Excel (live formulas), slides, SVG charts and an HTML page.

  python finman.py dcf MSFT [--set growth_y1=0.12 --set terminal_g=0.02] [--out DIR] [--formats xlsx,pptx,svg,html]
  python finman.py lbo CROX | comps MSFT --peers AAPL,GOOGL,META | three MSFT | ratios MSFT | analyze MSFT | score MSFT | earnings MSFT
  python finman.py budget household.json           a personal monthly budget (format: references/budget.md)
  python finman.py all MSFT --peers AAPL,GOOGL     every company model in one go
  python finman.py audit FILE.xlsx                 recalculate in LibreOffice, list formula errors and the Checks sheet
  python finman.py render FILE.xlsx|FILE.pptx      PDF + PNG pages (LibreOffice + pdftoppm) to look at the file
  python finman.py fetch MSFT                      the normalized facts as JSON
  python finman.py selftest                        offline check of every model and writer (no network)

Data: SEC EDGAR company facts (10-K and 10-Q), Yahoo chart prices, FRED DGS10. Set SEC_USER_AGENT="name email" (SEC asks for it).
Needs Python 3.10+, openpyxl and python-pptx. LibreOffice is optional (audit, render).
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import fin_models as F  # noqa: E402
import fin_report as RP  # noqa: E402
import fin_svg as S  # noqa: E402

KINDS = ("dcf", "lbo", "comps", "three", "ratios", "analyze", "score", "earnings")


def parse_sets(items):
    out = {}
    for it in items or []:
        k, _, v = it.partition("=")
        out[k.strip()] = float(v)
    return out


def write_outputs(rep, stem, out, formats, xlsx_fn=None):
    out.mkdir(parents=True, exist_ok=True)
    files = []
    if "xlsx" in formats and xlsx_fn:
        p, cells = xlsx_fn(str(out / f"{stem}.xlsx"))
        files.append(p)
        (out / f"{stem}.cells.json").write_text(json.dumps(cells, indent=1), encoding="utf-8")
    if "pptx" in formats:
        import fin_pptx
        files.append(fin_pptx.deck(rep, str(out / f"{stem}.pptx")))
    if "svg" in formats:
        for i, c in enumerate(rep.get("charts", []), start=1):
            p = out / f"{stem}-chart{i}.svg"
            p.write_text(S.chart_svg(c), encoding="utf-8")
            files.append(str(p))
    if "html" in formats:
        p = out / f"{stem}.html"
        p.write_text(S.page(rep), encoding="utf-8")
        files.append(str(p))
    (out / f"{stem}.report.json").write_text(json.dumps(rep, indent=1, default=str), encoding="utf-8")
    return files


def run_company(kind, f, args, peers=None, cache=None):
    """One model for one company: returns (report, xlsx builder)."""
    import fin_xlsx as X
    cache = cache if cache is not None else {}
    o = parse_sets(args.set)
    if kind in ("dcf", "score", "comps") and "dcf" not in cache:
        a = F.dcf_inputs(f, o if kind == "dcf" else None)
        cache["dcf"] = (a, F.dcf(a), F.dcf_grid(a))
    if kind == "dcf":
        a, r, g = cache["dcf"]
        return RP.dcf_report(f, a, r, g), lambda p: X.dcf_book(f, a, p)
    if kind == "lbo":
        a = F.lbo_inputs(f, o)
        return RP.lbo_report(f, a, F.lbo(a), F.lbo_grid(a)), lambda p: X.lbo_book(f, a, p)
    if kind == "comps":
        c = F.comps(f, peers)
        g = cache["dcf"][2][2]
        inner = [g[i][j] for i in (1, 2, 3) for j in (1, 2, 3)]
        return RP.comps_report(c, (min(inner), max(inner))), lambda p: X.comps_book(c, p)
    if kind == "three":
        a = F.three_inputs(f, o)
        m = F.three_statement(f, a)
        return RP.three_report(f, a, m), lambda p: X.three_book(f, a, m, p)
    if kind == "ratios":
        return RP.ratios_report(f, F.ratios(f)), lambda p: X.ratios_book(f, p)
    if kind == "analyze":
        fl = F.red_flags(f)
        return RP.analysis_report(f, F.ratios(f), fl, F.altman_z(f), F.piotroski(f)), lambda p: X.analysis_book(f, fl, p)
    if kind == "score":
        v = cache["dcf"][1]["value_per_share"]
        sc = F.scorecard(f, v)
        return RP.score_report(f, sc, v), lambda p: X.score_book(f, sc, v, p)
    if kind == "earnings":
        e = F.earnings(f)
        return RP.earnings_report(f, e), lambda p: X.earnings_book(f, e, p)
    raise SystemExit(f"unknown model: {kind}")


def soffice():
    for c in (shutil.which("soffice"), shutil.which("libreoffice"), r"C:\Program Files\LibreOffice\program\soffice.exe",
              "/Applications/LibreOffice.app/Contents/MacOS/soffice"):
        if c and Path(c).exists():
            return c
    return None


def recalc(path, outdir):
    """LibreOffice opens the workbook, computes every formula and saves the cached values next to them."""
    exe = soffice()
    if not exe:
        raise SystemExit("LibreOffice not found: install it, or open the file in Excel (it recalculates on open).")
    prof = Path(tempfile.mkdtemp(prefix="lo-"))
    subprocess.run([exe, f"-env:UserInstallation=file:///{prof.as_posix().lstrip('/')}", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", str(outdir), str(path)], check=True, capture_output=True, timeout=180)
    return Path(outdir) / Path(path).name


def audit(path, expect=None):
    """Formula errors, the Checks sheet, and (optional) cells that must equal the Python numbers. Returns (ok, lines)."""
    from openpyxl import load_workbook
    tmp = Path(tempfile.mkdtemp(prefix="finaudit-"))
    rc = recalc(path, tmp)
    wb = load_workbook(rc, data_only=True)
    raw = load_workbook(path)
    errors, formulas = [], 0
    for ws in raw.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas += 1
                    v = wb[ws.title][c.coordinate].value
                    if isinstance(v, str) and v.startswith("#") or v is None:
                        errors.append(f"{ws.title}!{c.coordinate} = {v} ({c.value[:60]})")
    lines = [f"{formulas} formulas, {len(errors)} errors"] + errors[:20]
    ok = not errors
    if "Checks" in wb.sheetnames:
        ws = wb["Checks"]
        for r in range(4, ws.max_row + 1):
            name, val = ws[f"A{r}"].value, ws[f"B{r}"].value
            if name:
                lines.append(f"check {'PASS' if val is True else 'FAIL'}: {name}")
                ok &= val is True
    for ref, want in (expect or {}).items():
        sh, cell = ref.split("!")
        got = wb[sh.strip("'")][cell].value
        same = isinstance(got, (int, float)) and abs(got - want) <= max(1e-6, abs(want) * 0.005)
        lines.append(f"cell {'PASS' if same else 'FAIL'}: {ref} = {got} (Python {want})")
        ok &= same
    return ok, lines


def render(path, outdir=None):
    exe = soffice()
    if not exe:
        raise SystemExit("LibreOffice not found")
    outdir = Path(outdir or Path(path).parent / "render" / Path(path).suffix.lstrip("."))
    outdir.mkdir(parents=True, exist_ok=True)
    prof = Path(tempfile.mkdtemp(prefix="lo-"))
    subprocess.run([exe, f"-env:UserInstallation=file:///{prof.as_posix().lstrip('/')}", "--headless", "--convert-to", "pdf", "--outdir",
                    str(outdir), str(path)], check=True, capture_output=True, timeout=180)
    pdf = outdir / (Path(path).stem + ".pdf")
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-png", "-r", "70", str(pdf), str(outdir / Path(path).stem)], check=True)
    return pdf


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("cmd", choices=KINDS + ("all", "budget", "fetch", "audit", "render", "selftest"))
    ap.add_argument("target", nargs="?")
    ap.add_argument("--peers", default="")
    ap.add_argument("--set", action="append", help="override an input: name=value (fractions for percentages)")
    ap.add_argument("--out", default="finman-out")
    ap.add_argument("--formats", default="xlsx,pptx,svg,html")
    ap.add_argument("--years", type=int, default=5)
    args = ap.parse_args(argv)
    out, formats = Path(args.out), set(args.formats.split(","))
    if args.cmd == "selftest":
        import fin_selftest
        return fin_selftest.run()
    if args.cmd == "audit":
        ok, lines = audit(args.target)
        print("\n".join(lines))
        print("PASS" if ok else "FAIL")
        return 0 if ok else 1
    if args.cmd == "render":
        print(render(args.target))
        return 0
    if args.cmd == "budget":
        import fin_xlsx as X
        b = F.budget(json.loads(Path(args.target).read_text(encoding="utf-8")))
        files = write_outputs(RP.budget_report(b), "budget", out, formats, lambda p: X.budget_book(b, p))
        print("\n".join(files))
        return 0
    import fin_data as D
    f = D.company(args.target, args.years)
    if args.cmd == "fetch":
        print(json.dumps(f, indent=1, default=str))
        return 0
    peers = [D.company(p, args.years) for p in args.peers.split(",") if p] if args.peers else []
    kinds = KINDS if args.cmd == "all" else (args.cmd,)
    if "comps" in kinds and not peers:
        if args.cmd == "comps":
            raise SystemExit("comps needs --peers A,B,C")
        kinds = tuple(k for k in kinds if k != "comps")
    cache = {}
    for kind in kinds:
        rep, xl = run_company(kind, f, args, peers, cache)
        files = write_outputs(rep, f"{f['ticker'].lower()}-{kind}", out, formats, xl)
        print(f"{kind}: {rep['hero']['value']} {rep['hero']['label']}")
        print("  " + "\n  ".join(files))
    return 0


if __name__ == "__main__":
    sys.exit(main())
